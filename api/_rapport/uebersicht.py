"""Blatt "Übersicht": Kopfband, 6 Kennzahl-Kacheln und 4 Diagramme auf einer A4-Seite quer."""

from datetime import date
from typing import Any

from openpyxl.chart._chart import ChartBase
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from . import stil
from .zellen import schreibe_text, schreibe_zahl

SPALTEN = 12  # A bis L; jede Kachel ist zwei Spalten breit
SPALTEN_BREITE = 11.5
LETZTE = get_column_letter(SPALTEN)
KACHEL_ZEILE = 5
DIAGRAMM_ZEILEN = (9, 24)  # Anker der beiden Diagrammreihen
DRUCK_ENDE = 38


def _band(ws: Any, zeile: int, von: int, bis: int, farbe: str) -> None:
    # Füllung auf jede Zelle, sonst bleiben verbundene Bereiche teilweise weiss.
    for s in range(von, bis + 1):
        ws.cell(row=zeile, column=s).fill = stil.fuellung(farbe)


def _kopf(ws: Any, zeitraum: str, erstellt: date) -> None:
    _band(ws, 1, 1, SPALTEN, stil.PETROL)
    _band(ws, 2, 1, SPALTEN, stil.NAVY)
    _band(ws, 3, 1, SPALTEN, stil.KORALLE)
    ws.merge_cells(f"A1:{LETZTE}1")
    titel = schreibe_text(ws, 1, 1, "immoheart · Jahresrapport")
    titel.font = stil.schrift(20, fett=True, farbe=stil.WEISS)
    titel.alignment = stil.LINKS_MITTE

    ws.merge_cells("A2:H2")
    zr = schreibe_text(ws, 2, 1, zeitraum)
    zr.font = stil.schrift(12, farbe=stil.WEISS)
    zr.alignment = stil.LINKS_MITTE
    ws.merge_cells(f"I2:{LETZTE}2")
    am = schreibe_text(ws, 2, 9, f"Erstellt am {erstellt.strftime('%d.%m.%Y')}")
    am.font = stil.schrift(10, farbe=stil.WEISS)
    am.alignment = Alignment(horizontal="right", vertical="center", indent=1)

    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 4  # schmale Akzentlinie in Koralle
    ws.row_dimensions[4].height = 10


def _kachel(ws: Any, nr: int, k: dict[str, Any]) -> None:
    s = 2 * nr + 1
    z = KACHEL_ZEILE
    for zeile in range(z, z + 3):
        for spalte in (s, s + 1):
            zelle = ws.cell(row=zeile, column=spalte)
            zelle.fill = stil.fuellung(stil.KACHEL_GRUND)
            zelle.border = stil.KACHEL_RAHMEN
        ws.merge_cells(start_row=zeile, start_column=s, end_row=zeile, end_column=s + 1)

    titel = schreibe_text(ws, z, s, k["titel"])
    titel.font = stil.schrift(10, fett=True, farbe=stil.PETROL)
    titel.alignment = stil.MITTE

    if k["wert"] is None:
        # Ohne Grundlage (z. B. keine neuen Anfragen) keine 0 % vortäuschen.
        wert = schreibe_text(ws, z + 1, s, "–")
    else:
        wert = schreibe_zahl(ws, z + 1, s, k["wert"])
        wert.number_format = stil.FORMAT_PROZENT if k["format"] == "prozent" else stil.FORMAT_ZAHL
    wert.font = stil.schrift(24, fett=True, farbe=stil.NAVY)
    wert.alignment = stil.MITTE

    text = schreibe_text(ws, z + 2, s, k["text"])
    text.font = stil.schrift(8.5, farbe=stil.TINTE_2)
    text.alignment = stil.MITTE


def baue_uebersicht(ws: Any, modell: dict[str, Any], diagramme: list[ChartBase], erstellt: date) -> None:
    ws.title = "Übersicht"
    ws.sheet_view.showGridLines = False
    for s in range(1, SPALTEN + 1):
        ws.column_dimensions[get_column_letter(s)].width = SPALTEN_BREITE

    _kopf(ws, modell["zeitraum"]["label"], erstellt)
    for nr, kachel in enumerate(modell["kennzahlen"]):
        _kachel(ws, nr, kachel)
    ws.row_dimensions[KACHEL_ZEILE].height = 30
    ws.row_dimensions[KACHEL_ZEILE + 1].height = 38
    ws.row_dimensions[KACHEL_ZEILE + 2].height = 34

    # 2 × 2 Raster: links Spalte A, rechts Spalte G.
    anker = [f"{sp}{z}" for z in DIAGRAMM_ZEILEN for sp in ("A", "G")]
    for diagramm, ort in zip(diagramme, anker):
        ws.add_chart(diagramm, ort)

    stil.seite_quer(ws, eine_seite_hoch=True)
    ws.print_area = f"A1:{LETZTE}{DRUCK_ENDE}"
    ws.print_options.horizontalCentered = True
