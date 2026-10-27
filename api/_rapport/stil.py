"""Markenfarben, Schriften und Zahlenformate des Rapports an einer Stelle."""

from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PETROL = "065A82"
NAVY = "21295C"
KORALLE = "FF6B6B"  # Akzent, nur sparsam (Trennlinie, Vermittlungen, KI-Hinweis)
HELLBLAU = "9CC5D6"
KACHEL_GRUND = "EEF4F7"
TINTE = "1F2933"
TINTE_2 = "5B6670"
WEISS = "FFFFFF"

SCHRIFT = "Calibri"

FORMAT_ZAHL = "#'##0"
FORMAT_CHF = "\"CHF \"#'##0"
FORMAT_DATUM = "dd.mm.yyyy"
FORMAT_PROZENT = "0%"
FORMAT_TAGE = "0.0"

TABELLEN_STIL = "TableStyleMedium2"


def fuellung(farbe: str) -> PatternFill:
    return PatternFill("solid", start_color=farbe, end_color=farbe)


def schrift(groesse: float = 11, fett: bool = False, farbe: str = TINTE, kursiv: bool = False) -> Font:
    return Font(name=SCHRIFT, size=groesse, bold=fett, italic=kursiv, color=farbe)


LINKS_OBEN_UMBRUCH = Alignment(horizontal="left", vertical="top", wrap_text=True)
LINKS_MITTE = Alignment(horizontal="left", vertical="center", indent=1)
MITTE = Alignment(horizontal="center", vertical="center", wrap_text=True)

_DUENN = Side(style="thin", color="D5DEE3")
KACHEL_RAHMEN = Border(left=_DUENN, right=_DUENN, top=_DUENN, bottom=_DUENN)


def seite_quer(ws: Any, eine_seite_hoch: bool) -> None:
    """A4 quer und auf Seitenbreite skaliert; die Übersicht zusätzlich auf eine Seite."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    # 0 = beliebig viele Seiten hoch; lange Datenblätter dürfen umbrechen.
    ws.page_setup.fitToHeight = 1 if eine_seite_hoch else 0
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
