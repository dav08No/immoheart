"""Baut die Excel-Arbeitsmappe aus dem geprüften Modell (Spike)."""

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

from .modell import QUELLEN, pruefe_modell

QUELLEN_TITEL = {"mail": "Mail", "website": "Website", "manuell": "Manuell"}


def _schreibe_daten(wb: Workbook, monate: list[dict[str, Any]]) -> Any:
    # Diagrammdaten liegen auf einem versteckten Blatt, damit "Übersicht"
    # druckfertig bleibt und die Diagramme trotzdem echte Zellbezüge haben.
    daten = wb.create_sheet("_daten")
    daten.sheet_state = "hidden"
    daten.append(["Monat", *(QUELLEN_TITEL[q] for q in QUELLEN)])
    for m in monate:
        daten.append([m["label"], *(m[q] for q in QUELLEN)])
    return daten


def _anfragen_diagramm(daten: Any, anzahl_monate: int) -> BarChart:
    diagramm = BarChart()
    diagramm.type = "col"
    diagramm.grouping = "stacked"
    # Ohne overlap=100 zeichnet Excel gestapelte Säulen versetzt.
    diagramm.overlap = 100
    diagramm.title = "Anfragen pro Monat nach Quelle"
    diagramm.y_axis.title = "Anfragen"
    letzte_zeile = max(anzahl_monate + 1, 2)
    werte = Reference(daten, min_col=2, max_col=1 + len(QUELLEN), min_row=1, max_row=letzte_zeile)
    kategorien = Reference(daten, min_col=1, min_row=2, max_row=letzte_zeile)
    diagramm.add_data(werte, titles_from_data=True)
    diagramm.set_categories(kategorien)
    diagramm.width = 24
    diagramm.height = 10
    return diagramm


def baue_arbeitsmappe(modell: dict[str, Any]) -> bytes:
    """Reine Funktion: Modell rein, .xlsx-Bytes raus. Prüft selbst nochmals."""
    sauber = pruefe_modell(modell)
    monate = sauber["diagramme"]["anfragenProMonat"]

    wb = Workbook()
    uebersicht = wb.active
    uebersicht.title = "Übersicht"
    uebersicht["A1"] = f"immoheart Jahresrapport · {sauber['zeitraum']['label']}"
    uebersicht["A1"].font = Font(size=16, bold=True, color="1F4E5F")

    daten = _schreibe_daten(wb, monate)
    uebersicht.add_chart(_anfragen_diagramm(daten, len(monate)), "A3")

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()
