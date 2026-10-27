"""Datenblätter als Excel-Tabellen mit Filter, fester Kopfzeile und Formaten."""

from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import stil
from .zellen import schreibe_text, schreibe_zahl

# Blattname und Tabellenname je Modellschlüssel; Tabellennamen sind eindeutig,
# ohne Leerzeichen und kollidieren nicht mit Zellbezügen wie "A1".
BLAETTER = {
    "vermittlungen": ("Vermittlungen", "tbl_Vermittlungen"),
    "angebote": ("Angebote", "tbl_Angebote"),
    "treffer": ("Treffer", "tbl_Treffer"),
    "anfragen": ("Anfragen", "tbl_Anfragen"),
    "objekte": ("Objekte", "tbl_Objekte"),
}
KEINE_EINTRAEGE = "Keine Einträge"
FORMATE = {"zahl": stil.FORMAT_ZAHL, "chf": stil.FORMAT_CHF, "datum": stil.FORMAT_DATUM}
MIN_BREITE = 10
MAX_BREITE = 45
# Für die Breite reicht eine Stichprobe; spart bei 20000 Zeilen Zeit.
BREITE_STICHPROBE = 500


def _breite(spalte: dict[str, str], zeilen: list[list[Any]], index: int) -> float:
    if spalte["typ"] == "datum":
        laenge = 10
    elif spalte["typ"] == "text":
        laenge = max((len(z[index]) for z in zeilen[:BREITE_STICHPROBE] if z[index]), default=0)
    else:
        laenge = 12 if spalte["typ"] == "chf" else 8
    # +4 Platz für den Filterpfeil der Kopfzeile.
    return float(min(MAX_BREITE, max(MIN_BREITE, laenge + 2, len(spalte["titel"]) + 4)))


def _zeile(ws: Any, nr: int, typen: list[str], werte: list[Any]) -> None:
    for s, (typ, wert) in enumerate(zip(typen, werte), start=1):
        if wert is None:
            continue
        if typ == "text":
            schreibe_text(ws, nr, s, wert)
        else:
            # Datum (date) und Zahlen sind geprüft; nur das Format unterscheidet sie.
            zelle = schreibe_zahl(ws, nr, s, wert)
            zelle.number_format = FORMATE[typ]


def baue_tabellenblatt(wb: Any, schluessel: str, tabelle: dict[str, Any]) -> Any:
    blatt_name, tabellen_name = BLAETTER[schluessel]
    ws = wb.create_sheet(blatt_name)
    spalten = tabelle["spalten"]
    zeilen = tabelle["zeilen"]
    typen = [s["typ"] for s in spalten]

    for s, spalte in enumerate(spalten, start=1):
        # Kopfzeilen-Farben liefert der Tabellenstil.
        schreibe_text(ws, 1, s, spalte["titel"])
        ws.column_dimensions[get_column_letter(s)].width = _breite(spalte, zeilen, s - 1)

    for nr, werte in enumerate(zeilen, start=2):
        _zeile(ws, nr, typen, werte)
    if not zeilen:
        # Excel-Tabellen brauchen mindestens eine Datenzeile, sonst meldet Excel einen Fehler.
        schreibe_text(ws, 2, 1, KEINE_EINTRAEGE).font = stil.schrift(11, kursiv=True, farbe=stil.TINTE_2)

    letzte_zeile = max(len(zeilen), 1) + 1
    bereich = f"A1:{get_column_letter(len(spalten))}{letzte_zeile}"
    t = Table(displayName=tabellen_name, ref=bereich)
    t.tableStyleInfo = TableStyleInfo(name=stil.TABELLEN_STIL, showRowStripes=True)
    ws.add_table(t)

    ws.freeze_panes = "A2"
    stil.seite_quer(ws, eine_seite_hoch=False)
    ws.print_title_rows = "1:1"  # Kopfzeile auf jeder gedruckten Seite
    return ws
