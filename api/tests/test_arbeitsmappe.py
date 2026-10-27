"""Tests für den Aufbau der Arbeitsmappe aus dem Fixture und seinen Varianten."""

import time
from datetime import date, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from _rapport import baue_arbeitsmappe
from conftest import KI_TEXT, leeres_modell, mappe, modell

BLAETTER = ["Übersicht", "Zusammenfassung", "Vermittlungen", "Angebote", "Treffer", "Anfragen", "Objekte",
            "Hinweise", "_daten"]
DATENBLAETTER = {"Vermittlungen": "vermittlungen", "Angebote": "angebote", "Treffer": "treffer",
                 "Anfragen": "anfragen", "Objekte": "objekte"}


def _texte(ws: Any) -> list[str]:
    return [c.value for zeile in ws.iter_rows() for c in zeile if isinstance(c.value, str)]


def test_blaetter_in_reihenfolge_und_daten_versteckt() -> None:
    wb = mappe(modell())
    assert wb.sheetnames == BLAETTER
    assert wb["_daten"].sheet_state == "hidden"
    assert all(wb[n].sheet_state == "visible" for n in BLAETTER[:-1])


def test_uebersicht_kopf_kacheln_und_diagramme() -> None:
    wb = load_workbook(BytesIO(baue_arbeitsmappe(modell(), heute=date(2026, 10, 27))))
    ws = wb["Übersicht"]
    assert ws["A1"].value == "immoheart · Jahresrapport"
    assert ws["A2"].value == "27.10.2025 – 27.10.2026"
    assert ws["I2"].value == "Erstellt am 27.10.2026"
    assert len(ws._charts) == 4
    # 6 Kacheln à zwei Spalten: Titel in Zeile 5, Zahl in Zeile 6.
    assert [ws.cell(5, 2 * i + 1).value for i in range(6)][0] == "Neue Anfragen"
    assert ws["A6"].value == 3 and ws["A6"].number_format == "#'##0"
    assert ws["K6"].value == pytest.approx(1 / 3) and ws["K6"].number_format == "0%"
    assert not ws.sheet_view.showGridLines


def test_quote_null_zeigt_strich() -> None:
    m = modell()
    m["kennzahlen"][5]["wert"] = None
    assert mappe(m)["Übersicht"]["K6"].value == "–"


def test_diagrammarten() -> None:
    charts = mappe(modell())["Übersicht"]._charts
    assert charts[0].grouping == "stacked"
    assert charts[2].type == "bar" and charts[3].type == "bar"
    assert all(c.legend.position == "b" for c in charts)


def test_seitenlayout_quer_eine_seite() -> None:
    ws = mappe(modell())["Übersicht"]
    assert ws.page_setup.orientation == "landscape"
    assert int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A4)
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 1
    assert ws.print_area


def test_tabellen_zeilenzahl_filter_und_kopfzeile() -> None:
    m = modell()
    wb = mappe(m)
    for blatt, schluessel in DATENBLAETTER.items():
        ws = wb[blatt]
        tabelle = next(iter(ws.tables.values()))
        anzahl = len(m["tabellen"][schluessel]["zeilen"])
        assert ws.max_row == anzahl + 1
        assert tabelle.ref.endswith(str(anzahl + 1))
        assert tabelle.autoFilter is not None
        assert tabelle.tableStyleInfo.name == "TableStyleMedium2"
        assert ws.freeze_panes == "A2"
        assert ws["A1"].value == m["tabellen"][schluessel]["spalten"][0]["titel"]


def test_tabellennamen_eindeutig_ohne_leerzeichen() -> None:
    wb = mappe(modell())
    namen = [n for b in DATENBLAETTER for n in wb[b].tables]
    assert len(set(namen)) == 5 and all(" " not in n for n in namen)


def test_datum_und_zahlenformate() -> None:
    ws = mappe(modell())["Vermittlungen"]
    assert ws["A2"].is_date and ws["A2"].value == datetime(2026, 1, 20)
    assert ws["A2"].number_format == "dd.mm.yyyy"
    assert ws["D2"].value == 300 and ws["D2"].number_format == "#'##0"
    assert ws["E2"].value == 210 and ws["E2"].number_format == '"CHF "#\'##0'


def test_null_zelle_bleibt_leer() -> None:
    ws = mappe(modell())["Treffer"]
    assert ws["F3"].value is None  # archiviertAm null


@pytest.mark.parametrize("boese", ["=1+1", '=HYPERLINK("http://x","y")', "+1", "-1", "@a", "\t=1"])
def test_formel_injection_in_tabelle_und_ki_text(boese: str) -> None:
    m = modell()
    m["tabellen"]["treffer"]["zeilen"][0]["firma"] = boese
    # Einzeilige Felder lehnen Tabs ab, deshalb dort ohne führenden Tab.
    m["tabellen"]["treffer"]["spalten"][0]["titel"] = boese.strip()
    m["zusammenfassung"] = {"text": {**KI_TEXT["text"], "rueckblick": boese}}
    m["diagramme"]["nutzungen"][0]["label"] = boese.strip()
    wb = mappe(m)
    for zelle, wert in ((wb["Treffer"]["A2"], boese), (wb["Treffer"]["A1"], boese.strip())):
        assert zelle.data_type == "s" and zelle.value == wert and zelle.quotePrefix
    ki = [c for c in wb["Zusammenfassung"]["A"] if c.value == boese]
    assert ki and ki[0].data_type == "s" and ki[0].quotePrefix
    assert wb["_daten"]["M2"].data_type == "s"


def test_ki_text() -> None:
    m = modell()
    m["zusammenfassung"] = KI_TEXT
    texte = _texte(mappe(m)["Zusammenfassung"])
    assert "Von KI verfasst auf Basis der Kennzahlen – bitte prüfen." in texte
    assert {"Jahresrückblick", "Trends", "Auffälligkeiten", "Ausblick"} <= set(texte)
    assert KI_TEXT["text"]["ausblick"] in texte


def test_ki_fehler() -> None:
    m = modell()
    m["zusammenfassung"] = {"fehler": "Tageslimit erreicht"}
    texte = _texte(mappe(m)["Zusammenfassung"])
    assert "KI-Zusammenfassung nicht verfügbar (Tageslimit erreicht)." in texte
    assert "Trends" not in texte


@pytest.mark.parametrize("fehlt", [True, False])
def test_ki_nicht_erstellt(fehlt: bool) -> None:
    m = modell()
    if not fehlt:
        m["zusammenfassung"] = None
    assert "KI-Zusammenfassung nicht erstellt." in _texte(mappe(m)["Zusammenfassung"])


def test_medianwerte() -> None:
    ws = mappe(modell())["Zusammenfassung"]
    werte = {ws.cell(c.row, 1).value: ws.cell(c.row, 2).value for c in ws["B"] if c.value is not None}
    assert werte == {"Tage bis Erstangebot (Median)": 2.5, "Tage bis Abschluss (Median)": 78}


def test_hinweise_als_liste() -> None:
    m = modell()
    texte = _texte(mappe(m)["Hinweise"])
    assert f"• {m['hinweise'][-1]}" in texte
    assert sum(t.startswith("• ") for t in texte) == len(m["hinweise"])


def test_leeres_modell_ergibt_gueltige_datei() -> None:
    wb = mappe(leeres_modell())
    assert wb.sheetnames == BLAETTER
    assert len(wb["Übersicht"]._charts) == 4
    for blatt in DATENBLAETTER:
        ws = wb[blatt]
        # Platzhalterzeile, weil Excel-Tabellen mindestens eine Datenzeile brauchen.
        assert ws["A2"].value == "Keine Einträge"
        assert next(iter(ws.tables.values())).ref.endswith("2")
    assert wb["_daten"]["A2"].value == "Keine Daten"


def test_5000_treffer_unter_30_sekunden() -> None:
    m = modell()
    vorlage = m["tabellen"]["treffer"]["zeilen"][0]
    m["tabellen"]["treffer"]["zeilen"] = [{**vorlage, "score": i % 100, "firma": f"Firma {i}"} for i in range(5000)]
    start = time.perf_counter()
    daten = baue_arbeitsmappe(m)
    dauer = time.perf_counter() - start
    assert dauer < 30
    ws = load_workbook(BytesIO(daten))["Treffer"]
    assert ws.max_row == 5001
    assert next(iter(ws.tables.values())).ref == "A1:G5001"
