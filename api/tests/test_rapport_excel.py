"""Tests für die Rapport-Funktion: reine Funktion und Handler mit gefälschtem Request."""

import json
import os
import sys
from io import BytesIO

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import rapport_excel  # noqa: E402
from _rapport import ModellFehler, baue_arbeitsmappe  # noqa: E402

SCHLUESSEL = "test-schluessel"

MODELL = {
    "zeitraum": {"label": "27.10.2025 – 27.10.2026"},
    "diagramme": {
        "anfragenProMonat": [
            {"label": "Okt 25", "mail": 1, "website": 2, "manuell": 0},
            {"label": "Nov 25", "mail": 3, "website": 0, "manuell": 1},
        ]
    },
}


class _FalscherSocket:
    """Ersetzt die Netzwerkverbindung: liefert den Request, sammelt die Antwort."""

    def __init__(self, roh: bytes) -> None:
        self._ein = BytesIO(roh)
        self.aus = BytesIO()

    def makefile(self, modus: str, *_args: object, **_kwargs: object) -> BytesIO:
        return self._ein if "r" in modus else self.aus

    def sendall(self, daten: bytes) -> None:
        self.aus.write(daten)


def _anfrage(body: bytes, headers: dict[str, str], methode: str = "POST") -> tuple[int, bytes]:
    kopf = f"{methode} /api/rapport_excel HTTP/1.1\r\nHost: test\r\n"
    kopf += "".join(f"{k}: {v}\r\n" for k, v in headers.items())
    sock = _FalscherSocket(kopf.encode("latin-1") + b"\r\n" + body)
    rapport_excel.handler(sock, ("127.0.0.1", 0), None)
    antwort = sock.aus.getvalue()
    status = int(antwort.split(b" ", 2)[1])
    return status, antwort.split(b"\r\n\r\n", 1)[1]


def _post(modell: object, schluessel: str | None = SCHLUESSEL) -> tuple[int, bytes]:
    body = json.dumps(modell).encode("utf-8")
    headers = {"Content-Length": str(len(body))}
    if schluessel is not None:
        headers["x-rapport-schluessel"] = schluessel
    return _anfrage(body, headers)


@pytest.fixture(autouse=True)
def _env(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("RAPPORT_SCHLUESSEL", SCHLUESSEL)


def test_gueltiges_modell_ergibt_xlsx_mit_diagramm() -> None:
    wb = load_workbook(BytesIO(baue_arbeitsmappe(MODELL)))
    assert wb.sheetnames == ["Übersicht", "_daten"]
    assert "27.10.2025 – 27.10.2026" in wb["Übersicht"]["A1"].value
    assert wb["_daten"].sheet_state == "hidden"
    assert wb["_daten"].max_row == 3
    assert len(wb["Übersicht"]._charts) == 1


def test_leere_monatsliste_ergibt_gueltige_datei() -> None:
    leer = {"zeitraum": {"label": "x"}, "diagramme": {"anfragenProMonat": []}}
    wb = load_workbook(BytesIO(baue_arbeitsmappe(leer)))
    assert "Übersicht" in wb.sheetnames


def test_handler_liefert_xlsx() -> None:
    status, body = _post(MODELL)
    assert status == 200
    assert body[:2] == b"PK"  # .xlsx ist ein ZIP-Archiv


@pytest.mark.parametrize("schluessel", [None, "falsch", ""])
def test_falscher_oder_fehlender_schluessel_401(schluessel: str | None) -> None:
    assert _post(MODELL, schluessel)[0] == 401


def test_ohne_env_schluessel_500(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("RAPPORT_SCHLUESSEL")
    assert _post(MODELL)[0] == 500


def test_ueber_5_mb_413() -> None:
    headers = {"Content-Length": str(5 * 1024 * 1024 + 1), "x-rapport-schluessel": SCHLUESSEL}
    assert _anfrage(b"", headers)[0] == 413


def test_kaputtes_json_400() -> None:
    body = b'{"zeitraum": '
    headers = {"Content-Length": str(len(body)), "x-rapport-schluessel": SCHLUESSEL}
    assert _anfrage(body, headers)[0] == 400


def test_get_405() -> None:
    assert _anfrage(b"", {}, "GET")[0] == 405


@pytest.mark.parametrize(
    "kaputt",
    [
        [],
        {"zeitraum": {"label": 5}, "diagramme": {"anfragenProMonat": []}},
        {"zeitraum": {"label": "x"}, "diagramme": {"anfragenProMonat": {}}},
        {"zeitraum": {"label": "x"}, "diagramme": {"anfragenProMonat": [], "extra": 1}},
        {**MODELL, "email": "a@b.ch"},
        {"zeitraum": {"label": "x"}, "diagramme": {"anfragenProMonat": [
            {"label": "Okt 25", "mail": True, "website": 0, "manuell": 0}]}},
        {"zeitraum": {"label": "x"}, "diagramme": {"anfragenProMonat": [
            {"label": "Okt 25", "mail": -1, "website": 0, "manuell": 0}]}},
        {"zeitraum": {"label": "x"}, "diagramme": {"anfragenProMonat": [
            {"label": "Okt 25", "mail": 1.5, "website": 0, "manuell": 0}]}},
    ],
)
def test_unerwartete_typen_400(kaputt: object) -> None:
    status, body = _post(kaputt)
    assert status == 400
    assert b"a@b.ch" not in body  # keine Datenteile in der Fehlerantwort
    with pytest.raises(ModellFehler):
        baue_arbeitsmappe(kaputt)  # type: ignore[arg-type]


@pytest.mark.parametrize("label", ["=1+1", '=HYPERLINK("http://x","y")', "+1", "-1", "@a"])
def test_formel_label_bleibt_text(label: str) -> None:
    modell = {"zeitraum": {"label": label}, "diagramme": {"anfragenProMonat": [
        {"label": label, "mail": 1, "website": 0, "manuell": 0}]}}
    wb = load_workbook(BytesIO(baue_arbeitsmappe(modell)))
    zelle = wb["_daten"]["A2"]
    assert zelle.data_type == "s" and zelle.value == label
    assert zelle.quotePrefix
    assert wb["Übersicht"]["A1"].data_type == "s"


def test_steuerzeichen_im_label_400() -> None:
    modell = {"zeitraum": {"label": "a\x00b"}, "diagramme": {"anfragenProMonat": []}}
    status, body = _post(modell)
    assert status == 400
    assert b"zeitraum.label" in body


def test_nicht_ascii_schluessel_401() -> None:
    assert _post(MODELL, "schlüssel-ä")[0] == 401


def test_tiefe_verschachtelung_400() -> None:
    body = b"[" * 100_000 + b"]" * 100_000
    headers = {"Content-Length": str(len(body)), "x-rapport-schluessel": SCHLUESSEL}
    assert _anfrage(body, headers)[0] == 400


@pytest.mark.parametrize("laenge", ["+5", "5 5", "1_0", "²", "-1"])
def test_ungueltige_content_length_411(laenge: str) -> None:
    headers = {"Content-Length": laenge, "x-rapport-schluessel": SCHLUESSEL}
    assert _anfrage(b"", headers)[0] == 411


def test_unerwarteter_fehler_500_ohne_daten(monkeypatch: pytest.MonkeyPatch) -> None:
    def kaputt(_modell: object) -> bytes:
        raise RuntimeError("geheim")

    monkeypatch.setattr(rapport_excel, "baue_arbeitsmappe", kaputt)
    status, body = _post(MODELL)
    assert status == 500
    assert json.loads(body) == {"fehler": "Interner Fehler"}
